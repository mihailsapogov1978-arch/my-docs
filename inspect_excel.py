#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Диагностика: показывает первые 5 строк Excel файла"""

import warnings
from openpyxl import load_workbook

warnings.filterwarnings("ignore", message="Workbook contains no default style")

filepath = 'docs/plans/2027/PI2027.xlsx'
wb = load_workbook(filepath, data_only=True)
ws = wb.worksheets[0]

print(f"Лист: {ws.title}")
print(f"Всего строк: {ws.max_row}\n")

print("=== ПЕРВЫЕ 5 СТРОК ДАННЫХ ===")
for row in range(1, 6):
    print(f"\nСтрока {row}:")
    for col in range(1, 11):
        cell = ws.cell(row=row, column=col)
        if cell.value is not None:
            print(f"  Колонка {col}: {repr(cell.value)}")

wb.close()