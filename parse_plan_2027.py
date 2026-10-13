#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Парсит PI2027.xlsx (все годы на одном листе)
Группирует по годам и генерирует один Markdown файл с разделами
"""

import os
import re
import warnings
from datetime import datetime
from openpyxl import load_workbook

warnings.filterwarnings("ignore", message="Workbook contains no default style")

EXCEL_FILE = 'docs/plans/PI2027.xlsx'
OUTPUT_MD = 'docs/plans/plans.md'

RKS_LINKS = {
    2027: 'https://rks.yanao.ru/application/?format=storeLink&storeId=bbBfdai',
    2026: 'https://rks.yanao.ru/application/?format=storeLink&storeId=YOUR_2026_STORE_ID',
    2025: 'https://rks.yanao.ru/application/?format=storeLink&storeId=YOUR_2025_STORE_ID',
}

CURRENT_YEAR = 2026


def parse_excel(filepath):
    """Читает Excel и возвращает список записей"""
    if not os.path.exists(filepath):
        print(f'❌ Файл не найден: {filepath}')
        return None
    
    wb = load_workbook(filepath, data_only=True)
    ws = wb.worksheets[0]
    
    records = []
    
    for row in range(2, ws.max_row + 1):
        status = ws.cell(row=row, column=1).value
        event_num = ws.cell(row=row, column=2).value
        customer = ws.cell(row=row, column=4).value
        
        if not event_num:
            continue
        
        if customer != 'ГКУ ЯНАО "ЦБ ОГВ ЯНАО"':
            continue
        
        year_match = re.match(r'^(\d{4})\.', str(event_num).strip())
        if not year_match:
            continue
        year = int(year_match.group(1))
            
        records.append({
            'year': year,
            'status': status or '—',
            'event_num': str(event_num).strip(),
            'name': ws.cell(row=row, column=5).value or '—',
            'amount': ws.cell(row=row, column=6).value,
        })
    
    wb.close()
    return records


def fmt_amount(v):
    """Форматирует сумму"""
    if v is None or v == '':
        return '—'
    try:
        return f'{float(v):,.0f}'.replace(',', ' ')
    except (ValueError, TypeError):
        return str(v)


def build_table_html(records, year):
    """Генерирует HTML-таблицу для конкретного года (без колонки «Согласовано»)."""
    
    total_amount = 0
    for r in records:
        try:
            total_amount += float(r['amount']) if r['amount'] and str(r['amount']).replace('.','').isdigit() else 0
        except:
            pass
    
    rows_html = ''
    for i, r in enumerate(records, start=1):
        amount = fmt_amount(r['amount'])
        name = r['name'].replace('\n', '<br>')
        
        rows_html += '<tr>'
        rows_html += '<td class="pt-num">' + str(i) + '</td>'
        rows_html += '<td class="pt-name">' + name + '</td>'
        rows_html += '<td class="pt-amount">' + amount + ' ₽</td>'
        rows_html += '</tr>'

    rks_link = RKS_LINKS.get(year, '#')
    
    if year > CURRENT_YEAR:
        link_text = '[Смотреть предварительный ПИ' + str(year) + ' на rks.yanao.ru](' + rks_link + ')' if rks_link != '#' else ''
    else:
        link_text = '[Смотреть ПИ' + str(year) + ' на rks.yanao.ru](' + rks_link + ')' if rks_link != '#' else ''
    
    total_fmt = fmt_amount(total_amount)
    
    html = link_text + '\n\n'
    html += '<div class="plan-table-wrap">\n'
    html += '<table class="plan-table">\n'
    html += '  <thead>\n'
    html += '    <tr>\n'
    html += '      <th class="pt-num">№</th>\n'
    html += '      <th class="pt-name">Наименование мероприятия</th>\n'
    html += '      <th class="pt-amount">Объем расходов</th>\n'
    html += '    </tr>\n'
    html += '  </thead>\n'
    html += '  <tbody>\n'
    html += rows_html
    html += '    <tr class="pt-total">\n'
    html += '      <td colspan="2" class="pt-total-label">ИТОГО:</td>\n'
    html += '      <td class="pt-amount pt-total-amount">' + total_fmt + ' ₽</td>\n'
    html += '    </tr>\n'
    html += '  </tbody>\n'
    html += '</table>\n'
    html += '</div>\n'
    
    return html


def main():
    print('=' * 60)
    print('Генерация планов информатизации (все годы)')
    print('=' * 60)
    
    print('\n📖 Чтение ' + EXCEL_FILE + '...')
    records = parse_excel(EXCEL_FILE)
    
    if records is None:
        return
    
    print('📊 Всего найдено записей (ГКУ ЯНАО "ЦБ ОГВ ЯНАО"): ' + str(len(records)))
    
    if len(records) == 0:
        print('⚠️  Данные не найдены.')
        return
    
    plans_by_year = {}
    for r in records:
        year = r['year']
        if year not in plans_by_year:
            plans_by_year[year] = []
        plans_by_year[year].append(r)
    
    print('📅 Найдено годов: ' + str(sorted(plans_by_year.keys(), reverse=True)))
    
    md_content = '''<style>
.plan-table-wrap {
    margin: 12px 0 24px 0;
    width: 100%;
}
.plan-table {
    width: 100%;
    border-collapse: collapse;
    font-size: 14px;
    table-layout: fixed;
}
.plan-table thead th {
    background: #f5f7fa;
    color: #000;
    text-align: left;
    padding: 10px 12px;
    font-weight: 600;
    font-size: 13px;
    border-bottom: 1px solid #e2e8f0;
}
.plan-table tbody td {
    padding: 10px 12px;
    border-bottom: 1px solid #eef2f6;
    vertical-align: top;
}
.plan-table tbody tr:last-child td { border-bottom: none; }
.plan-table .pt-num    { width: 44px; text-align: center; color: #718096; white-space: nowrap; }
.plan-table .pt-name {
    /* ширина по умолчанию auto — займёт всё свободное место */
}
.plan-table .pt-amount {
    width: 150px;
    text-align: right;
    white-space: nowrap;
    font-weight: 500;
}
.plan-table .pt-total {
    background: #e8f4f8;
    font-size: 15px;
}
.plan-table .pt-total-label {
    text-align: right;
    font-weight: 600;
    padding: 10px 12px;
}
.plan-table .pt-total-amount {
    font-weight: 700;
}
@media (max-width: 720px) {
    .plan-table { font-size: 13px; table-layout: auto; }
    .plan-table .pt-num    { width: 28px; }
    .plan-table .pt-amount { width: 110px; }
}
</style>

# План закупок по мероприятиям информатизации
'''
    
    for year in sorted(plans_by_year.keys(), reverse=True):
        year_records = plans_by_year[year]
        print('\n Генерация раздела для ' + str(year) + ' года (' + str(len(year_records)) + ' записей)...')
        
        md_content += '## ' + str(year) + ' год\n\n'
        md_content += build_table_html(year_records, year)
        md_content += '\n---\n\n'
    
    os.makedirs(os.path.dirname(OUTPUT_MD), exist_ok=True)
    with open(OUTPUT_MD, 'w', encoding='utf-8') as f:
        f.write(md_content)
    
    print('\n✅ Готово: ' + OUTPUT_MD)
    print('📑 Сгенерировано разделов: ' + str(len(plans_by_year)))


if __name__ == '__main__':
    main()