#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Собирает единый отчёт docs/reports/index.md из всех .xlsx в папке reports_src/
(или из конкретного файла, переданного аргументом).

Каждый месяц = один H2 → попадает в правый TOC.
Аккордеон сотрудников:
    СВЁРНУТО:  ФИО   Звонков N   Задач M   Заявок K   Резолюций R
    РАЗВЁРНУТО: полный список задач

Сотрудники без задач в отчёт не попадают.
Счётчики заявок/резолюций берутся из столбца J листа.
"""
import os
import re
import sys
import glob
import html
from datetime import datetime

from openpyxl import load_workbook

MONTHS_RU = {
    'январь': 1, 'февраль': 2, 'март': 3, 'апрель': 4, 'май': 5, 'июнь': 6,
    'июль': 7, 'август': 8, 'сентябрь': 9, 'октябрь': 10, 'ноябрь': 11, 'декабрь': 12,
}
MONTHS_RU_NOM = {
    1: 'Январь', 2: 'Февраль', 3: 'Март', 4: 'Апрель', 5: 'Май', 6: 'Июнь',
    7: 'Июль', 8: 'Август', 9: 'Сентябрь', 10: 'Октябрь', 11: 'Ноябрь', 12: 'Декабрь',
}

SHEET_NAME = 'Количественная аналитика'
SRC_DIR = 'reports_src'
OUT_PATH = os.path.join('docs', 'reports', 'index.md')
COUNTERS_COL = 10  # столбец J (1-based)

SKIP_TASKS = {
    'текущие задачи', 'исполнитель', 'работа с заявками',
    'текущая задача', 'задачи', 'мероприятие',
}


def esc(v):
    if v is None:
        return ''
    return html.escape(str(v)).replace('\n', '<br>')


def detect_month(xlsx_path):
    base = os.path.basename(xlsx_path).lower()
    m = None
    for name, num in MONTHS_RU.items():
        if name in base:
            m = num
            break
    y = re.search(r'(20\d{2})', base)
    year = int(y.group(1)) if y else datetime.now().year
    month = m or datetime.now().month
    return year, month


def find_month_row(ws, year, month):
    month_name = MONTHS_RU_NOM[month].lower()
    for r in range(1, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if not a:
            continue
        s = str(a).lower()
        if month_name in s and str(year) in s:
            return r
    last = None
    for r in range(1, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if not a:
            continue
        if month_name in str(a).lower():
            last = r
    return last


def extract_calls(ws, year, month):
    employees = []
    col = 3
    while True:
        v = ws.cell(row=2, column=col).value
        if not v:
            break
        name = str(v).strip()
        if name.lower().startswith('работа'):
            break
        employees.append((col, name))
        col += 1

    row_num = find_month_row(ws, year, month)
    if row_num is None:
        return []

    result = []
    for col, name in employees:
        v = ws.cell(row=row_num, column=col).value
        try:
            v = int(v) if v is not None else 0
        except (TypeError, ValueError):
            v = 0
        result.append((name, v))
    return result


def parse_counters(raw):
    """
    Примеры входных строк:
      'Исполнено резолюций - 11, в работе - 27, заявок в службу технической поддержки - 1'
      'Исполнено резолюций - 2, направлено заявок в службу технической поддержки - 3 (Смета)'
      'Направлено заявок в службу технической поддержки - 0'
      'Направлено заявок в службу технической поддержки - 7 заявок'
    """
    if not raw:
        return {}
    s = str(raw)
    counters = {}

    m = re.search(r'резолюц\w*\s*[-–—:]\s*(\d+)', s, re.IGNORECASE)
    if m:
        counters['resolutions'] = int(m.group(1))

    m = re.search(r'в\s+работе\s*[-–—:]\s*(\d+)', s, re.IGNORECASE)
    if m:
        counters['resolutions_inwork'] = int(m.group(1))

    m = re.search(r'заявок[^\d]*?[-–—:]\s*(\d+)', s, re.IGNORECASE)
    if m:
        counters['tickets'] = int(m.group(1))

    return counters


def extract_employees(ws):
    emp_row = None
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip().lower() == 'исполнитель':
            emp_row = r
            break
    if emp_row is None:
        return []

    employees = []
    r = emp_row + 1
    while r <= ws.max_row:
        name_cell = ws.cell(row=r, column=1).value
        name = str(name_cell).strip() if name_cell else ''

        if not name:
            r += 1
            continue

        if name.startswith('=') or name.lower() in SKIP_TASKS:
            r += 1
            continue

        raw_tasks = ws.cell(row=r, column=2).value
        tasks = []
        if raw_tasks:
            for line in str(raw_tasks).split('\n'):
                s = line.strip()
                if not s:
                    continue
                if s.lower() in SKIP_TASKS:
                    continue
                s = re.sub(r'^\d+[\.\)]\s*', '', s)
                s = re.sub(r'^[-–—]\s*', '', s)
                tasks.append(s)

        counters = {}
        raw_counters = ws.cell(row=r, column=COUNTERS_COL).value
        if raw_counters:
            counters = parse_counters(raw_counters)

        employees.append({
            'name': name,
            'tasks': tasks,
            'tickets': counters.get('tickets'),
            'resolutions': counters.get('resolutions'),
        })
        r += 1

    return employees


def render_employees(calls, employees):
    calls_map = {name: v for name, v in calls}

    employees = [e for e in employees if e['tasks']]
    if not employees:
        return ''

    total_calls = sum(calls_map.get(e['name'], 0) for e in employees)
    total_tasks = sum(len(e['tasks']) for e in employees)
    total_tickets = sum(e['tickets'] for e in employees if e.get('tickets') is not None)
    total_resolutions = sum(e['resolutions'] for e in employees if e.get('resolutions') is not None)

    rows = []
    for e in employees:
        name = e['name']
        calls_val = calls_map.get(name, 0)
        n_tasks = len(e['tasks'])

        chips = [
            f'<span class="chip calls">Звонков: <b>{calls_val}</b></span>',
            f'<span class="chip tasks">Задач: <b>{n_tasks}</b></span>',
        ]
        if e.get('tickets') is not None:
            chips.append(f'<span class="chip tickets">Заявок: <b>{e["tickets"]}</b></span>')
        if e.get('resolutions') is not None:
            chips.append(f'<span class="chip res">Резолюций: <b>{e["resolutions"]}</b></span>')

        chips_html = ''.join(chips)
        full_list = ''.join(f'<li>{esc(t)}</li>' for t in e['tasks'])

        rows.append(
            '<details class="emp-row">'
            '<summary class="emp-summary">'
            f'<span class="emp-name">{esc(name)}</span>'
            f'<span class="emp-chips">{chips_html}</span>'
            '</summary>'
            '<div class="emp-body">'
            f'<ul class="emp-tasks">{full_list}</ul>'
            '</div>'
            '</details>'
        )

        header = (
            '<div class="emp-total">'
            '<span class="emp-total-name">Отдел СРТП</span>'
            f'<span class="emp-total-cell">Сотрудников: <b>{len(employees)}</b></span>'
            f'<span class="emp-total-cell">Звонков: <b>{total_calls}</b></span>'
            f'<span class="emp-total-cell">Задач: <b>{total_tasks}</b></span>'
            f'<span class="emp-total-cell">Заявок: <b>{total_tickets}</b></span>'
            f'<span class="emp-total-cell">Резолюций: <b>{total_resolutions}</b></span>'
            '</div>'
        )

    return (
        '<div markdown="1">\n\n'
        + header + '\n\n'
        + '<div class="emp-list">\n'
        + '\n'.join(rows) + '\n'
        + '</div>\n\n'
        + '</div>'
    )


STYLES = '''<style>
/* ===== Единая сетка колонок для шапки и строк ===== */
.emp-total,
.md-typeset .emp-summary,
.emp-summary {
    display: grid !important;
    /*          ФИО    Сотрудников  Звонков  Задач   Заявок   Резолюций */
    grid-template-columns: 1fr  120px  130px  110px  110px  120px;
    align-items: center !important;
    gap: 0 !important;
    margin: 0 !important;
    padding: 0 !important;
}

/* ===== Шапка отдела ===== */
.emp-total {
    background: #ebf8ff;
    border: 1px solid #bee3f8;
    border-radius: 8px;
    margin: 10px 0 12px 0 !important;
    font-size: 13px;
    color: #2c5282;
}
.emp-total > span {
    padding: 10px 8px;
    white-space: nowrap;
}
.emp-total-name {
    padding-left: 14px !important;
    font-weight: 600;
    color: #2b6cb0;
}
.emp-total b { color: #2b6cb0; }

/* ===== Список сотрудников ===== */
.emp-list {
    margin: 16px 0;
    border: 1px solid #e2e8f0;
    border-radius: 10px;
    overflow: hidden;
    background: #fff;
}

.md-typeset .emp-row,
.emp-row {
    border: none !important;
    outline: none !important;
    box-shadow: none !important;
    border-bottom: 1px solid #eef2f6 !important;
    margin: 0 !important;
    padding: 0 !important;
}
.emp-row:last-child { border-bottom: none !important; }

/* Summary — тоже grid, те же колонки */
.md-typeset .emp-summary,
.emp-summary {
    cursor: pointer;
    list-style: none !important;
    font-size: 13px;
    line-height: 1.3;
    transition: background 0.15s;
    outline: none !important;
    box-shadow: none !important;
    border: none !important;
    padding-inline-start: 0 !important;
}

/* Все маркеры summary — убрать */
.md-typeset .emp-summary::-webkit-details-marker,
.emp-summary::-webkit-details-marker { display: none !important; }
.md-typeset .emp-summary::marker,
.emp-summary::marker { display: none !important; content: "" !important; }
.md-typeset .emp-summary::before,
.emp-summary::before { display: none !important; content: none !important; }

.emp-summary:hover { background: #f7fafc; }

/* Без синей рамки фокуса */
.md-typeset .emp-row:focus,
.md-typeset .emp-row:focus-visible,
.md-typeset .emp-row:focus-within,
.md-typeset .emp-summary:focus,
.md-typeset .emp-summary:focus-visible {
    outline: none !important;
    box-shadow: none !important;
    border-color: transparent !important;
}

/* ===== Ячейки строки ===== */
.emp-summary > span {
    padding: 10px 8px;
    vertical-align: middle;
    white-space: nowrap;
}

.md-typeset .emp-name,
.emp-name {
    font-weight: 700;
    color: #2c3e50;
    text-transform: uppercase;
    letter-spacing: 0.02em;
    padding-left: 14px !important;
    text-align: left;
}

/* Чипы — «раскрываем» в те же ячейки grid-сетки */
.emp-chips { display: contents !important; }

.md-typeset .chip,
.chip {
    padding: 10px 8px !important;
    font-size: 12px;
    line-height: 1.3;
    white-space: nowrap !important;
    background: transparent !important;
    border-radius: 0 !important;
    text-align: left !important;
    width: auto !important;
}
.chip b { font-weight: 700; }

.chip.calls  { color: #2b6cb0; }
.chip.tasks  { color: #2c7a7b; }
.chip.tickets{ color: #744210; }
.chip.res    { color: #6b46c1; }

/* ===== Тело раскрытой карточки ===== */
.md-typeset .emp-body,
.emp-body {
    padding: 4px 14px 12px 14px !important;
    background: #fafbfc;
    font-size: 12.5px;
    color: #2d3748;
}
.emp-tasks {
    margin: 0;
    padding-left: 18px;
    line-height: 1.5;
}
.emp-tasks li { margin-bottom: 4px; }

/* Узкие экраны */
@media (max-width: 900px) {
    .emp-total,
    .md-typeset .emp-summary,
    .emp-summary {
        grid-template-columns: 1fr !important;
        gap: 6px !important;
    }
    .emp-total > span,
    .emp-summary > span { padding: 4px 14px; }
    .emp-chips {
        display: flex !important;
        flex-wrap: wrap;
        gap: 8px;
        padding-left: 14px;
    }
    .chip { padding: 2px 8px !important; }
}
</style>'''


def build_month_section(xlsx_path):
    year, month = detect_month(xlsx_path)

    try:
        wb = load_workbook(xlsx_path, data_only=True)
    except Exception as e:
        print(f'  ⚠️  Не удалось открыть {xlsx_path}: {e}')
        return None

    if SHEET_NAME not in wb.sheetnames:
        print(f'  ⚠️  В {xlsx_path} нет листа "{SHEET_NAME}", пропускаю')
        return None

    ws = wb[SHEET_NAME]
    calls = extract_calls(ws, year, month)
    employees = extract_employees(ws)
    month_label = f'{MONTHS_RU_NOM[month]} {year}'

    parts = []
    parts.append(f'## {month_label}')
    parts.append('')

    html_block = render_employees(calls, employees)
    if not html_block:
        parts.append(f'*Данные за {month_label.lower()} не найдены.*')
        parts.append('')
        return year, month, '\n'.join(parts)

    parts.append(html_block)
    parts.append('')

    return year, month, '\n'.join(parts)


def main():
    if len(sys.argv) > 1:
        files = [sys.argv[1]]
    else:
        files = sorted(glob.glob(os.path.join(SRC_DIR, '*.xlsx')))

    if not files:
        print('Нет .xlsx-файлов. Положите их в reports_src/ или укажите файл аргументом.')
        sys.exit(1)

    sections = []
    for f in files:
        print(f'Обрабатываю: {f}')
        res = build_month_section(f)
        if res:
            sections.append(res)

    if not sections:
        print('Не удалось собрать ни одного отчёта.')
        sys.exit(1)

    sections.sort(key=lambda t: (t[0], t[1]), reverse=True)

    out = [STYLES, '', '# Отчёты отдела СРТП', '']
    for _, _, text in sections:
        out.append(text)
        out.append('')

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    with open(OUT_PATH, 'w', encoding='utf-8') as fh:
        fh.write('\n'.join(out))

    print(f'✅ Готово: {OUT_PATH}')
    print(f'   Отчётов: {len(sections)}')


if __name__ == '__main__':
    main()